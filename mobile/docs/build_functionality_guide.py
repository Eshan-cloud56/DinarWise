from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "DINARWISE_FUNCTIONALITY_AND_PUBLISHING_GUIDE.md"
OUTPUT = ROOT / "DinarWise_Functionality_and_Android_Publishing_Guide.docx"

TEAL = RGBColor(0x0D, 0x7A, 0x67)
DARK = RGBColor(0x18, 0x27, 0x23)
MUTED = RGBColor(0x5A, 0x68, 0x64)
PALE = "E7F3EF"
LIGHT = "F4F8F6"


def set_font(run, size=None, bold=None, color=None, italic=None, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def create_numbering_instance(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1.")
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "432")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "432")
    indent.set(qn("w:hanging"), "252")
    ppr.extend([tabs, indent])
    level.extend([start, fmt, text, suffix, ppr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    num.append(reference)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.extend([level, number])
    ppr.append(num_pr)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0D7A67")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text):
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|https?://\S+)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_font(paragraph.add_run(text[cursor:match.start()]), size=10.5, color=DARK)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=9.5, color=TEAL, name="Menlo")
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=10.5, bold=True, color=DARK)
        else:
            clean = token.rstrip(".,;)")
            trailing = token[len(clean):]
            add_hyperlink(paragraph, clean, clean)
            if trailing:
                set_font(paragraph.add_run(trailing), size=10.5, color=DARK)
        cursor = match.end()
    if cursor < len(text):
        set_font(paragraph.add_run(text[cursor:]), size=10.5, color=DARK)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.14

    heading_tokens = {
        "Heading 1": (17, 14, 6),
        "Heading 2": (13.5, 11, 4),
        "Heading 3": (11.5, 8, 3),
    }
    for name, (size, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = TEAL if name != "Heading 3" else DARK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Aptos"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Inches(0.28)
    bullet.paragraph_format.first_line_indent = Inches(-0.16)
    bullet.paragraph_format.space_after = Pt(3)

    numbered = doc.styles["List Number"]
    numbered.font.name = "Aptos"
    numbered.font.size = Pt(10.5)
    numbered.paragraph_format.left_indent = Inches(0.30)
    numbered.paragraph_format.first_line_indent = Inches(-0.18)
    numbered.paragraph_format.space_after = Pt(3)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("DINAR WISE  |  FUNCTIONALITY & ANDROID RELEASE GUIDE")
    set_font(run, size=8, bold=True, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Dinar Wise • Android • 4 August 2026    |    ")
    set_font(run, size=8, color=MUTED)
    add_field(p, "PAGE")


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("PRODUCT & RELEASE REFERENCE")
    set_font(run, size=10, bold=True, color=TEAL)
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Dinar Wise")
    set_font(run, size=34, bold=True, color=DARK, name="Aptos Display")
    title.paragraph_format.space_after = Pt(3)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Expense AI Manager")
    set_font(run, size=18, color=TEAL, name="Aptos Display")
    subtitle.paragraph_format.space_after = Pt(10)

    descriptor = doc.add_paragraph()
    descriptor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = descriptor.add_run(
        "Functionality, architecture, privacy, analytics, and Android publishing guide"
    )
    set_font(run, size=11.5, color=MUTED)
    descriptor.paragraph_format.space_after = Pt(30)

    table = doc.add_table(rows=4, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.6)
    metadata = [
        ("Application ID", "com.sl.dinarwise.expensemanager"),
        ("Platform", "Android"),
        ("Storage", "Offline Drift/SQLite • Schema 2"),
        ("Telemetry", "Optional Firebase • Consent controlled"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(row.cells[0], PALE)
        set_cell_fill(row.cells[1], LIGHT)
        lp = row.cells[0].paragraphs[0]
        set_font(lp.add_run(label), size=9.5, bold=True, color=TEAL)
        vp = row.cells[1].paragraphs[0]
        set_font(vp.add_run(value), size=9.5, color=DARK)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(28)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "Core financial records stay on the device. Android Firebase telemetry is optional."
    )
    set_font(run, size=10, italic=True, color=MUTED)

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(40)
    set_font(date.add_run("Prepared 4 August 2026"), size=9.5, color=MUTED)
    doc.add_page_break()


def add_document_notice(doc):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.65)
    cell = table.cell(0, 0)
    set_cell_fill(cell, PALE)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    p = cell.paragraphs[0]
    set_font(p.add_run("Release status: "), size=10, bold=True, color=TEAL)
    set_font(
        p.add_run(
            "the latest APK is verified for device testing, but Play Store publication still requires a protected production upload key and a signed release AAB."
        ),
        size=10,
        color=DARK,
    )
    doc.add_paragraph()


def render_markdown(doc, markdown):
    in_code = False
    code_lines = []
    skipped_header = 0
    active_numbering = None
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.22)
                p.paragraph_format.right_indent = Inches(0.22)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(7)
                p_pr = p._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), LIGHT)
                p_pr.append(shd)
                for index, code in enumerate(code_lines):
                    if index:
                        p.add_run().add_break()
                    run = p.add_run(code)
                    set_font(run, size=8.7, color=DARK, name="Menlo")
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line:
            active_numbering = None
            continue
        if skipped_header < 2 and line.startswith("#"):
            skipped_header += 1
            continue
        if line == "---":
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            active_numbering = None
            level = min(len(heading.group(1)), 3)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, heading.group(2))
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            active_numbering = None
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            if active_numbering is None:
                active_numbering = create_numbering_instance(doc)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            apply_numbering(p, active_numbering)
            add_inline(p, numbered.group(1))
            continue
        active_numbering = None
        p = doc.add_paragraph()
        add_inline(p, line)


def main():
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_header_footer(doc.sections[0])
    add_cover(doc)
    add_document_notice(doc)
    render_markdown(doc, markdown)
    props = doc.core_properties
    props.title = "Dinar Wise Functionality and Android Publishing Guide"
    props.subject = "Product functionality, privacy, Firebase, and Google Play publication"
    props.author = "Dinar Wise Project"
    props.keywords = "DinarWise, Flutter, Android, Firebase, Google Play, finance"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
